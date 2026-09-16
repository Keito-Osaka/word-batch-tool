import argparse, json, os, sys, tempfile, traceback, zipfile
import pandas as pd
import core

def write_json(stream,value): stream.buffer.write((json.dumps(value,ensure_ascii=False)+"\n").encode("utf-8")); stream.buffer.flush()
def emit_progress(stage,label,current=0,total=0,percent=None,detail="",state="running"):
    write_json(sys.stdout,{"type":"progress","stage":stage,"label":label,"current":current,"total":total,"percent":percent,"detail":detail,"state":state})
def read_request():
    raw=sys.stdin.buffer.read(); return json.loads(raw.decode("utf-8")) if raw else {}
def row_dict(row): return {str(k):"" if v is None else str(v) for k,v in row.to_dict().items()}
def validate_file(path,label):
    if not path: raise ValueError(f"{label}が指定されていません。")
    if not os.path.isfile(path): raise FileNotFoundError(f"{label}が見つかりません。\n\n{path}")
def load_raw(path):
    validate_file(path,"置換データファイル"); ext=os.path.splitext(path)[1].lower()
    if ext==".csv":
        last=None
        for enc in ("utf-8-sig","utf-8","cp932"):
            try: df=pd.read_csv(path,dtype=str,encoding=enc); break
            except UnicodeDecodeError as e: last=e
        else: raise ValueError("CSVの文字コードを判定できませんでした。") from last
        df=df.dropna(axis=1,how="all"); df.columns=core.deduplicate_columns([str(c).strip() for c in df.columns])
    elif ext in (".xlsx",".xlsm",".xls"):
        engine="openpyxl" if ext in (".xlsx",".xlsm") else "xlrd"
        df=pd.read_excel(path,sheet_name=0,header=0,dtype=str,engine=engine).dropna(axis=1,how="all")
        df.columns=core.deduplicate_columns([core.normalize_excel_column_name(c) for c in df.columns])
    else: raise ValueError("対応していない形式です。CSVまたはExcelを指定してください。")
    df=df.loc[:,[c for c in df.columns if str(c).strip()]]
    if len(df.columns)==0: raise ValueError("使用できる列がありません。")
    return df
def exclusion_reason(row, mode, columns):
    empty = [str(c) for c in columns if c in row.index and core.is_empty_cell_for_row_exclusion(row[c])]
    if mode == "selected_columns_all_empty":
        return "、".join(f"「{c}」" for c in columns) + "がすべて空欄"
    if mode == "selected_columns_any_empty":
        return "、".join(f"「{c}」" for c in empty) + "が空欄"
    check = list(row.index[1:]) if mode in ("any_empty_except_first", "all_empty_except_first") else list(row.index)
    empty = [str(c) for c in check if core.is_empty_cell_for_row_exclusion(row[c])]
    names = "、".join(f"「{c}」" for c in empty[:3]) + ("ほか" if len(empty) > 3 else "")
    return names + ("がすべて空欄" if mode == "all_empty_except_first" else "が空欄")

def exclusion_summary(mode, columns, exclude_examples):
    parts = ["1列目に「例」を含む行を除外" if exclude_examples else "記入例行も使用"]
    labels = {
        "none": "空欄による除外なし",
        "any_empty_except_first": "1列目以外に空欄があれば除外",
        "any_empty": "どこかに空欄があれば除外",
        "all_empty_except_first": "1列目以外がすべて空なら除外",
        "selected_columns_all_empty": "指定列がすべて空なら除外",
        "selected_columns_any_empty": "指定列のどれか1つでも空なら除外",
    }
    parts.append(labels.get(mode, mode))
    if mode.startswith("selected_columns") and columns:
        parts.append("対象列: " + "、".join(columns))
    return " / ".join(parts)

def inspect_data(req):
    df=load_raw(req.get("data_path","")); df=core.clean_dataframe_values(df); df=core.format_bank_account_columns(df)
    df=core.format_amount_columns(df,enabled=req.get("format_amount_with_comma",True),include_keywords=req.get("amount_include_keywords"),exclude_keywords=req.get("amount_exclude_keywords"))
    original=len(df); first=df.columns[0]; example_mask=df[first].astype(str).str.contains("例",na=False); examples=int(example_mask.sum())
    mode=req.get("row_exclude_mode",core.DEFAULT_ROW_EXCLUDE_MODE)
    columns=[c for c in req.get("row_exclude_columns",[]) if c in df.columns]
    exclude_examples=req.get("exclude_example_rows",True)
    included=[]; excluded=[]
    for idx,row in df.iterrows():
        item=row_dict(row)
        if exclude_examples and bool(example_mask.loc[idx]):
            excluded.append({"source_row_number":int(idx)+2,"row":item,"reason":"1列目に「例」を含むため除外"})
        elif core.should_exclude_row_by_mode(row,mode,req.get("row_exclude_target_column_number",1),columns):
            excluded.append({"source_row_number":int(idx)+2,"row":item,"reason":exclusion_reason(row,mode,columns)})
        else: included.append(item)
    return {"original_count":original,"example_count":examples,"excluded_count":len(excluded),"included_count":len(included),"columns":[str(c) for c in df.columns],"included_rows":included[:500],"excluded_rows":excluded[:500],"exclusion_summary":exclusion_summary(mode,columns,exclude_examples)}

def aggregate_path(template,folder,count,ext):
    stem=core.sanitize_filename_part(os.path.splitext(os.path.basename(template))[0]); base=f"{stem}_{count}件一式.{ext}"; return core.ensure_unique_path(os.path.join(folder,base))
def warmup():
    # 文書生成・Excel・PDF処理で使う主要モジュールを初回起動時に読み込む。
    import numpy
    import openpyxl
    import pypdf
    import docx
    return {"success":True,"python":sys.version.split()[0],"pandas":pd.__version__,"numpy":numpy.__version__}

def warmup_template(req):
    path=req.get("template_path","")
    validate_file(path,"テンプレートWord")
    # 保存や置換は行わず、構造を読み取ってOSとPythonのキャッシュを温める。
    document=core.Document(path)
    paragraph_count=len(document.paragraphs)
    table_count=len(document.tables)
    section_count=len(document.sections)
    for section in document.sections:
        _=len(section.header.paragraphs)+len(section.footer.paragraphs)
    fields=core.inspect_template_placeholders(path)
    word_fonts=core.inspect_effective_placeholder_fonts_with_word(path)
    return {"success":True,"paragraph_count":paragraph_count,"table_count":table_count,"section_count":section_count,**fields,"mixed_font_fields":word_fonts["mixed_font_fields"]}

def zip_files(folder,zip_path,total):
    files=[f for f in os.listdir(folder) if os.path.isfile(os.path.join(folder,f))]
    with zipfile.ZipFile(zip_path,"w",zipfile.ZIP_DEFLATED) as z:
        for i,name in enumerate(files,1): z.write(os.path.join(folder,name),arcname=name); emit_progress("zip","ZIPファイルにまとめています",i,total,92+int(7*i/max(total,1)),name)
def generate(req):
    template=req.get("template_path",""); data=req.get("data_path",""); output=req.get("output_path","")
    validate_file(template,"テンプレートWord"); validate_file(data,"置換データファイル")
    if not output: raise ValueError("出力先が指定されていません。")
    os.makedirs(output,exist_ok=True); emit_progress("prepare","置換データを準備しています",0,0,2)
    df=core.prepare_dataframe(data,format_amount_with_comma=req.get("format_amount_with_comma",True),amount_include_keywords=req.get("amount_include_keywords"),amount_exclude_keywords=req.get("amount_exclude_keywords"),row_exclude_mode=req.get("row_exclude_mode",core.DEFAULT_ROW_EXCLUDE_MODE),row_exclude_target_column_number=req.get("row_exclude_target_column_number",core.DEFAULT_ROW_EXCLUDE_TARGET_COLUMN_NUMBER),row_exclude_columns=req.get("row_exclude_columns",[]),exclude_example_rows=req.get("exclude_example_rows",True))
    word_fonts=core.inspect_effective_placeholder_fonts_with_word(template)
    core.set_word_effective_font_profiles(word_fonts["profiles"])
    common_values={str(k):"" if v is None else str(v) for k,v in req.get("common_values",{}).items()}
    template_fields=core.inspect_template_placeholders(template)
    missing=[name for name in template_fields["common_fields"] if not common_values.get(name,"").strip()]
    if missing: raise ValueError("共通項目に未入力があります。\n\n"+"\n".join(f"・{name}" for name in missing))
    count=len(df); emit_progress("prepare","置換データの準備が完了しました",count,count,8,state="done")
    keys=req.get("filename_keys",[]); serial=req.get("add_serial_number",True); digits=req.get("serial_digits",2)
    if not serial and not keys:
        raise ValueError("通し番号を付けない場合は、ファイル名に使用する列を1つ以上選択してください。")
    fmt=req.get("output_format","word"); method=req.get("output_method","folder"); fast=req.get("fast_pdf_split_enabled",True); actual=output
    def callback_for(stage,label,start,end):
        def cb(current,total): emit_progress(stage,label,current,total,start+int((end-start)*current/max(total,1)))
        return cb
    if fmt=="word" and method=="folder": core.generate_documents_from_template(template,df,keys,output,serial,digits,common_replacements=common_values,progress_callback=callback_for("word","個別Wordを作成しています",8,98))
    elif fmt=="word" and method=="merged": actual=aggregate_path(template,output,count,"docx"); core.generate_merged_document_from_template(template,df,actual,serial,digits,common_replacements=common_values,progress_callback=callback_for("word","結合Wordを作成しています",8,96)); emit_progress("save","結合Wordを保存しています",count,count,98)
    elif fmt=="word" and method=="zip":
        actual=aggregate_path(template,output,count,"zip")
        with tempfile.TemporaryDirectory() as tmp: core.generate_documents_from_template(template,df,keys,tmp,serial,digits,common_replacements=common_values,progress_callback=callback_for("word","ZIP用Wordを作成しています",8,90)); zip_files(tmp,actual,count)
    elif fmt=="pdf" and method=="folder":
        if fast:
            emit_progress("word","PDF変換用の結合Wordを作成しています",0,count,10)
            def fallback(pages,records): emit_progress("fallback","高速分割できないため通常方式へ切り替えます",0,records,None,state="notice")
            core.generate_pdf_documents_fast_or_fallback(template,df,keys,output,serial,digits,common_replacements=common_values,progress_callback=callback_for("word","PDF変換用の文書を作成しています",10,55),fallback_callback=fallback)
        else: core.generate_pdf_documents_from_template(template,df,keys,output,serial,digits,common_replacements=common_values,progress_callback=callback_for("pdf","個別PDFを作成しています",8,98))
    elif fmt=="pdf" and method=="merged": actual=aggregate_path(template,output,count,"pdf"); core.generate_merged_pdf_document_from_template(template,df,actual,serial,digits,common_replacements=common_values,progress_callback=callback_for("word","PDF変換用の結合Wordを作成しています",8,60)); emit_progress("pdf","Microsoft WordでPDFへ変換しています",count,count,92)
    elif fmt=="pdf" and method=="zip":
        actual=aggregate_path(template,output,count,"zip")
        with tempfile.TemporaryDirectory() as tmp:
            if fast:
                def fallback(pages,records): emit_progress("fallback","高速分割できないため通常方式へ切り替えます",0,records,None,state="notice")
                core.generate_pdf_documents_fast_or_fallback(template,df,keys,tmp,serial,digits,common_replacements=common_values,progress_callback=callback_for("word","ZIP用PDFを準備しています",8,82),fallback_callback=fallback)
            else: core.generate_pdf_documents_from_template(template,df,keys,tmp,serial,digits,common_replacements=common_values,progress_callback=callback_for("pdf","ZIP用PDFを作成しています",8,88))
            zip_files(tmp,actual,count)
    else: raise ValueError("出力形式または出力方法が正しくありません。")
    emit_progress("complete","処理が完了しました",count,count,100,state="done")
    return {"success":True,"output_path":actual,"output_directory":output,"generated_count":count,"format":fmt,"method":method}
def main():
    p=argparse.ArgumentParser(); p.add_argument("command",choices=["warmup","warmup-template","inspect","generate"]); a=p.parse_args()
    try:
        req=read_request(); result=warmup() if a.command=="warmup" else warmup_template(req) if a.command=="warmup-template" else inspect_data(req) if a.command=="inspect" else generate(req); write_json(sys.stdout,{"type":"result","data":result} if a.command=="generate" else result)
    except Exception as e: write_json(sys.stderr,{"error":str(e),"traceback":traceback.format_exc()}); sys.exit(1)
if __name__=="__main__":main()
