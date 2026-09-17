import os
import re
import tempfile
import zipfile
import unicodedata
import copy
import warnings

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DEFAULT_AMOUNT_INCLUDE_KEYWORDS = [
    "交付申請額", "交付決定額", "申請額", "金額", "費用", "参加費用", "支援金",
]
DEFAULT_AMOUNT_EXCLUDE_KEYWORDS = [
    "金融機関CD", "支店CD", "預金種別", "口座番号", "人数", "学年", "CD", "コード",
]
DEFAULT_ROW_EXCLUDE_MODE = "any_empty_except_first"
DEFAULT_ROW_EXCLUDE_TARGET_COLUMN_NUMBER = 1
DEFAULT_EXCLUDE_EXAMPLE_ROWS = True
# =====================================================
# DataFrame 読み込み・整形
# =====================================================
def deduplicate_columns(columns):
    result = []
    counts = {}
    for col in columns:
        base = str(col).strip()
        if base == "":
            result.append("")
            continue
        if base not in counts:
            counts[base] = 1
            result.append(base)
        else:
            counts[base] += 1
            result.append(f"{base}_{counts[base]}")
    return result


def normalize_excel_column_name(col):
    if col is None:
        return ""
    name = str(col).strip().replace("\r\n", "\n").replace("\r", "\n")
    compact = re.sub(r"\s+", "", name)
    replacements = [
        (("名前", "漢字"), "名前（漢字）"),
        (("名前", "カナ"), "名前（カナ）"),
        (("プログラム", "参加費用"), "プログラム参加費用"),
        (("本件以外", "支援金"), "本件以外の支援金"),
    ]
    for keys, value in replacements:
        if all(k in compact for k in keys):
            return value
    for key in ["金融機関CD", "支店CD", "預金種別", "口座番号", "口座名義人"]:
        if key in compact:
            return key
    return compact


def clean_cell_value(value):
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def clean_dataframe_values(df):
    df = df.copy()
    for col in df.columns:
        df[col] = df[col].apply(clean_cell_value)
    return df


def extract_digit_string(value):
    value = clean_cell_value(value)
    if value == "":
        return ""
    value = unicodedata.normalize("NFKC", value)
    value = value.replace(" ", "").replace("　", "")
    if re.fullmatch(r"\d+\.0", value):
        value = value[:-2]
    return re.sub(r"\D", "", value)


def zero_fill_code(value, length):
    digits = extract_digit_string(value)
    if digits == "":
        return ""
    return digits.zfill(length)


def format_account_number(value):
    digits = extract_digit_string(value)
    if digits == "":
        return ""
    if len(digits) == 8:
        digits = digits[:-1]
    return digits.zfill(7)


def format_bank_account_columns(df):
    df = df.copy()
    if "金融機関CD" in df.columns:
        df["金融機関CD"] = df["金融機関CD"].apply(lambda v: zero_fill_code(v, 4))
    if "支店CD" in df.columns:
        df["支店CD"] = df["支店CD"].apply(lambda v: zero_fill_code(v, 3))
    if "口座番号" in df.columns:
        df["口座番号"] = df["口座番号"].apply(format_account_number)
    return df


def _normalize_keyword_list(keywords):
    if keywords is None:
        return []
    return [re.sub(r"\s+", "", str(k).strip()) for k in keywords if str(k).strip()]


def is_amount_column(column_name, include_keywords=None, exclude_keywords=None):
    include_keywords = _normalize_keyword_list(
        DEFAULT_AMOUNT_INCLUDE_KEYWORDS if include_keywords is None else include_keywords
    )
    exclude_keywords = _normalize_keyword_list(
        DEFAULT_AMOUNT_EXCLUDE_KEYWORDS if exclude_keywords is None else exclude_keywords
    )
    name = re.sub(r"\s+", "", str(column_name).strip())
    if not name:
        return False
    if any(keyword in name for keyword in exclude_keywords):
        return False
    return any(keyword in name for keyword in include_keywords)


def format_amount_value(value):
    value = clean_cell_value(value)
    if value == "":
        return ""
    raw = unicodedata.normalize("NFKC", value).replace(",", "")
    raw = raw.replace(" ", "").replace("　", "")
    if re.fullmatch(r"-?\d+\.0", raw):
        raw = raw[:-2]
    if not re.fullmatch(r"-?\d+", raw):
        return value
    return f"{int(raw):,}"


def format_amount_columns(df, enabled=False, include_keywords=None, exclude_keywords=None):
    if not enabled:
        return df
    df = df.copy()
    for col in df.columns:
        if is_amount_column(col, include_keywords, exclude_keywords):
            df[col] = df[col].apply(format_amount_value)
    return df


# =====================================================
# 行除外ロジック
# =====================================================
def is_empty_cell_for_row_exclusion(value):
    if pd.isna(value):
        return True
    return str(value).strip() == ""


def _selected_values(row, row_exclude_columns=None, row_exclude_target_column_number=1):
    columns = [str(c) for c in (row_exclude_columns or []) if str(c) in row.index]
    if columns:
        return columns, [row[c] for c in columns]
    try:
        index = max(0, int(row_exclude_target_column_number) - 1)
    except Exception:
        index = 0
    if index < len(row.index):
        column = str(row.index[index])
        return [column], [row.iloc[index]]
    return [], []


def should_exclude_row_by_mode(row, row_exclude_mode=DEFAULT_ROW_EXCLUDE_MODE,
                             row_exclude_target_column_number=DEFAULT_ROW_EXCLUDE_TARGET_COLUMN_NUMBER,
                             row_exclude_columns=None):
    values = list(row.values)
    if row_exclude_mode == "none":
        return False
    if not values:
        return True
    other_values = values[1:]
    if row_exclude_mode == "any_empty_except_first":
        return any(is_empty_cell_for_row_exclusion(v) for v in other_values)
    if row_exclude_mode == "any_empty":
        return any(is_empty_cell_for_row_exclusion(v) for v in values)
    if row_exclude_mode == "all_empty_except_first":
        return all(is_empty_cell_for_row_exclusion(v) for v in other_values)
    columns, selected = _selected_values(row, row_exclude_columns, row_exclude_target_column_number)
    if not columns:
        return False
    if row_exclude_mode == "selected_columns_all_empty":
        return all(is_empty_cell_for_row_exclusion(v) for v in selected)
    if row_exclude_mode in ("selected_columns_any_empty", "selected_column_number_empty"):
        return any(is_empty_cell_for_row_exclusion(v) for v in selected)
    return False


def remove_example_and_empty_rows(df, row_exclude_mode=DEFAULT_ROW_EXCLUDE_MODE,
                                 row_exclude_target_column_number=DEFAULT_ROW_EXCLUDE_TARGET_COLUMN_NUMBER,
                                 row_exclude_columns=None, exclude_example_rows=DEFAULT_EXCLUDE_EXAMPLE_ROWS):
    if df.empty:
        return df
    result = df.copy()
    if exclude_example_rows:
        first_col = result.columns[0]
        result = result[~result[first_col].astype(str).str.contains("例", na=False)]
    result = result[~result.apply(lambda row: should_exclude_row_by_mode(
        row, row_exclude_mode, row_exclude_target_column_number, row_exclude_columns), axis=1)]
    return result.reset_index(drop=True)


def prepare_csv_dataframe(file_path, format_amount_with_comma=False, amount_include_keywords=None,
                          amount_exclude_keywords=None, row_exclude_mode=DEFAULT_ROW_EXCLUDE_MODE,
                          row_exclude_target_column_number=DEFAULT_ROW_EXCLUDE_TARGET_COLUMN_NUMBER, row_exclude_columns=None,
                          exclude_example_rows=DEFAULT_EXCLUDE_EXAMPLE_ROWS):
    encodings = ["utf-8-sig", "utf-8", "cp932"]
    last_error = None
    for enc in encodings:
        try:
            df = pd.read_csv(file_path, dtype=str, encoding=enc)
            break
        except Exception as e:
            last_error = e
    else:
        raise last_error
    df = df.dropna(axis=1, how="all")
    df.columns = deduplicate_columns([str(col).strip() for col in df.columns])
    df = df.loc[:, [col for col in df.columns if str(col).strip() != ""]]
    df = clean_dataframe_values(df)
    df = format_bank_account_columns(df)
    df = format_amount_columns(df, enabled=format_amount_with_comma,
                               include_keywords=amount_include_keywords,
                               exclude_keywords=amount_exclude_keywords)
    df = remove_example_and_empty_rows(df, row_exclude_mode=row_exclude_mode,
                                       row_exclude_target_column_number=row_exclude_target_column_number, row_exclude_columns=row_exclude_columns,
                                       exclude_example_rows=exclude_example_rows)
    if df.empty:
        raise ValueError("置換データがありません。CSVの内容を確認してください。")
    return df


def prepare_excel_dataframe(file_path, format_amount_with_comma=False, amount_include_keywords=None,
                           amount_exclude_keywords=None, row_exclude_mode=DEFAULT_ROW_EXCLUDE_MODE,
                           row_exclude_target_column_number=DEFAULT_ROW_EXCLUDE_TARGET_COLUMN_NUMBER, row_exclude_columns=None,
                           exclude_example_rows=DEFAULT_EXCLUDE_EXAMPLE_ROWS):
    ext = os.path.splitext(file_path)[1].lower()
    if ext in [".xlsx", ".xlsm"]:
        engine = "openpyxl"
    elif ext == ".xls":
        engine = "xlrd"
    else:
        raise ValueError("対応していないExcel形式です。")
    df = pd.read_excel(file_path, sheet_name=0, header=0, dtype=str, engine=engine)
    df = df.dropna(axis=1, how="all")
    df.columns = deduplicate_columns([normalize_excel_column_name(col) for col in df.columns])
    df = df.loc[:, [col for col in df.columns if str(col).strip() != ""]]
    df = clean_dataframe_values(df)
    df = format_bank_account_columns(df)
    df = format_amount_columns(df, enabled=format_amount_with_comma,
                               include_keywords=amount_include_keywords,
                               exclude_keywords=amount_exclude_keywords)
    df = remove_example_and_empty_rows(df, row_exclude_mode=row_exclude_mode,
                                       row_exclude_target_column_number=row_exclude_target_column_number, row_exclude_columns=row_exclude_columns,
                                       exclude_example_rows=exclude_example_rows)
    if df.empty:
        raise ValueError("置換データがありません。Excelの内容を確認してください。")
    return df


def prepare_dataframe(file_path, format_amount_with_comma=False, amount_include_keywords=None,
                     amount_exclude_keywords=None, row_exclude_mode=DEFAULT_ROW_EXCLUDE_MODE,
                     row_exclude_target_column_number=DEFAULT_ROW_EXCLUDE_TARGET_COLUMN_NUMBER, row_exclude_columns=None,
                     exclude_example_rows=DEFAULT_EXCLUDE_EXAMPLE_ROWS):
    if not file_path:
        raise ValueError("置換データファイルが指定されていません。")
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"置換データファイルが見つかりません: {file_path}")
    ext = os.path.splitext(file_path)[1].lower()
    kwargs = {
        "format_amount_with_comma": format_amount_with_comma,
        "amount_include_keywords": amount_include_keywords,
        "amount_exclude_keywords": amount_exclude_keywords,
        "row_exclude_mode": row_exclude_mode,
        "row_exclude_target_column_number": row_exclude_target_column_number,
        "row_exclude_columns": row_exclude_columns,
        "exclude_example_rows": exclude_example_rows,
    }
    if ext == ".csv":
        return prepare_csv_dataframe(file_path, **kwargs)
    if ext in [".xlsx", ".xlsm", ".xls"]:
        return prepare_excel_dataframe(file_path, **kwargs)
    raise ValueError("対応していないファイル形式です。CSVまたはExcelファイルを指定してください。")


# =====================================================
# Word置換・フォント維持
# =====================================================
_DEFAULT_FALLBACK_FONT = "ＭＳ 明朝"
_FONT_FALLBACK_WARNED = set()


def _script_for_char(char):
    if not char:
        return None
    code = ord(char)
    if 0x3040 <= code <= 0x30FF or 0x3400 <= code <= 0x4DBF or 0x4E00 <= code <= 0x9FFF:
        return "eastAsia"
    if 0xAC00 <= code <= 0xD7AF:
        return "eastAsia"
    if 0x0021 <= code <= 0x007E:
        return "ascii"
    return None


def _lookup_font_value(run, attribute_name):
    rPr = run._element.rPr
    if rPr is None or rPr.rFonts is None:
        return None
    value = rPr.rFonts.get(qn(f"w:{attribute_name}"))
    return value if value else None


def _detect_run_font(run, preferred_script=None):
    if run is None:
        return None

    rPr = run._element.rPr
    if rPr is not None and rPr.rFonts is not None:
        order = ("eastAsia", "ascii", "hAnsi", "cs")
        if preferred_script == "eastAsia":
            order = ("eastAsia", "ascii", "hAnsi", "cs")
        elif preferred_script == "ascii":
            order = ("ascii", "hAnsi", "eastAsia", "cs")
        elif preferred_script == "hAnsi":
            order = ("hAnsi", "ascii", "eastAsia", "cs")

        for attribute_name in order:
            value = _lookup_font_value(run, attribute_name)
            if value:
                return value

    if getattr(run, "font", None) is not None and run.font.name:
        font_name = run.font.name.strip()
        if font_name:
            return font_name

    return None


def get_run_preferred_font(run, default_font=_DEFAULT_FALLBACK_FONT, preferred_script=None):
    font_name = _detect_run_font(run, preferred_script=preferred_script)
    return font_name if font_name else default_font


def set_run_font_all(run, font_name):
    if not font_name:
        return

    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)

    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"), font_name)
    run.font.name = font_name


def copy_run_format(source_run, target_run):
    source_rPr = source_run._element.rPr
    target_rPr = target_run._element.rPr
    if target_rPr is not None:
        target_run._element.remove(target_rPr)
    if source_rPr is not None:
        target_run._element.insert(0, copy.deepcopy(source_rPr))


def _find_placeholder_source_run_index(start, end, char_to_run_index, original_text):
    for pos in range(start, end):
        if original_text[pos] not in "{}<>":
            return char_to_run_index[pos]
    return char_to_run_index[start] if start < len(char_to_run_index) else 0


def _find_placeholder_font(paragraph, start, end, char_to_run_index, original_text, field_name):
    first_run_index = None
    first_char = ""
    delimiter_run_indexes = []

    for pos in range(start, end):
        run_index = char_to_run_index[pos]
        ch = original_text[pos]
        if ch in "{}<>":
            if run_index not in delimiter_run_indexes:
                delimiter_run_indexes.append(run_index)
        elif first_run_index is None:
            first_run_index = run_index
            first_char = ch

    # 1. 先頭文字のフォントを優先
    if first_run_index is not None:
        font_name = _detect_run_font(
            paragraph.runs[first_run_index],
            preferred_script=_script_for_char(first_char),
        )
        if font_name:
            return font_name

    # 2. {{ / }} または << / >> のフォントを次に見る
    for run_index in delimiter_run_indexes:
        font_name = _detect_run_font(paragraph.runs[run_index])
        if font_name:
            return font_name

    # 3. 同じ段落の近い run を順に検索
    run_count = len(paragraph.runs)
    candidates = []

    before_index = char_to_run_index[start - 1] if start > 0 else None
    after_index = char_to_run_index[end] if end < len(char_to_run_index) else None
    for run_index in (before_index, after_index):
        if run_index is not None and run_index not in candidates:
            candidates.append(run_index)

    if first_run_index is not None:
        for distance in range(1, run_count + 1):
            for run_index in (first_run_index - distance, first_run_index + distance):
                if 0 <= run_index < run_count and run_index not in candidates:
                    candidates.append(run_index)

    for run_index in candidates:
        font_name = _detect_run_font(paragraph.runs[run_index])
        if font_name:
            return font_name

    # 4. どこにも取得できなければ警告付きで ＭＳ 明朝 を使用
    if field_name not in _FONT_FALLBACK_WARNED:
        _FONT_FALLBACK_WARNED.add(field_name)
        warnings.warn(
            "一部フォント情報を取得できなかったため、ＭＳ明朝を使用しました。"
            f" 対象: {field_name}",
            RuntimeWarning,
            stacklevel=3,
        )
    return _DEFAULT_FALLBACK_FONT


def replace_text_in_paragraph(paragraph, replacements):
    if not paragraph.runs:
        return

    original_text = "".join(run.text for run in paragraph.runs)
    if not original_text:
        return

    char_to_run_index = []
    for run_index, run in enumerate(paragraph.runs):
        char_to_run_index.extend([run_index] * len(run.text))

    matches = []
    replacement_groups = [("{{", "}}", replacements)]
    common_replacements = getattr(paragraph, "_word_batch_common_replacements", None)
    if common_replacements:
        replacement_groups.append(("<<", ">>", common_replacements))

    for opener, closer, values in replacement_groups:
        for key, value in values.items():
            placeholder = f"{opener}{key}{closer}"
            start = 0
            while True:
                pos = original_text.find(placeholder, start)
                if pos == -1:
                    break
                matches.append({
                    "start": pos,
                    "end": pos + len(placeholder),
                    "value": "" if value is None else str(value),
                    "opener": opener,
                    "closer": closer,
                })
                start = pos + len(placeholder)

    if not matches:
        return

    matches.sort(key=lambda x: x["start"])
    filtered = []
    last_end = -1
    for m in matches:
        if m["start"] >= last_end:
            filtered.append(m)
            last_end = m["end"]
    matches = filtered

    def add_original_segments(start, end):
        if start >= end:
            return
        segment_start = start
        while segment_start < end and segment_start < len(char_to_run_index):
            run_index = char_to_run_index[segment_start]
            segment_end = segment_start + 1
            while segment_end < end and segment_end < len(char_to_run_index) and char_to_run_index[segment_end] == run_index:
                segment_end += 1
            text = original_text[segment_start:segment_end]
            if text:
                segments.append({
                    "text": text,
                    "source_run_index": run_index,
                    "is_replacement": False,
                    "font_name": None,
                })
            segment_start = segment_end

    segments = []
    cursor = 0

    for m in matches:
        add_original_segments(cursor, m["start"])

        source_run_index = _find_placeholder_source_run_index(
            m["start"],
            m["end"],
            char_to_run_index,
            original_text,
        )
        font_name = _find_placeholder_font(
            paragraph,
            m["start"],
            m["end"],
            char_to_run_index,
            original_text,
            f"{m['opener']}{m['closer']}",
        )

        segments.append({
            "text": m["value"],
            "source_run_index": source_run_index,
            "is_replacement": True,
            "font_name": font_name,
        })
        cursor = m["end"]

    add_original_segments(cursor, len(original_text))

    existing_runs = list(paragraph.runs)
    for run in existing_runs:
        run.text = ""

    for i, segment in enumerate(segments):
        target_run = existing_runs[i] if i < len(existing_runs) else paragraph.add_run()

        source_index = segment["source_run_index"]
        if source_index >= len(existing_runs):
            source_index = 0

        source_run = existing_runs[source_index]
        copy_run_format(source_run, target_run)

        if segment["is_replacement"]:
            set_run_font_all(target_run, segment["font_name"])

        target_run.text = segment["text"]


def replace_text_in_table(table, replacements, common_replacements=None):
    for row in table.rows:
        for cell in row.cells:
            replace_text_in_document_part(cell, replacements, common_replacements)


def replace_text_in_document_part(part, replacements, common_replacements=None):
    for paragraph in part.paragraphs:
        paragraph._word_batch_common_replacements = common_replacements or {}
        replace_text_in_paragraph(paragraph, replacements)
    for table in part.tables:
        replace_text_in_table(table, replacements, common_replacements)


def replace_placeholders(doc, replacements, common_replacements=None):
    replace_text_in_document_part(doc, replacements, common_replacements)
    for section in doc.sections:
        replace_text_in_document_part(section.header, replacements, common_replacements)
        replace_text_in_document_part(section.footer, replacements, common_replacements)


def iter_document_text(doc):
    def walk(part):
        for paragraph in part.paragraphs:
            yield "".join(run.text for run in paragraph.runs)
        for table in part.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for section in doc.sections:
        yield from walk(section.header)
        yield from walk(section.footer)


def inspect_template_placeholders(template_path):
    doc = Document(template_path)
    text = "\n".join(iter_document_text(doc))
    row_fields = sorted(set(re.findall(r"\{\{([^{}\r\n]+)\}\}", text)))
    common_fields = sorted(set(re.findall(r"<<([^<>\r\n]+)>>", text)))
    malformed = []
    for line in text.splitlines():
        if ("<<" in line or ">>" in line) and not re.search(r"<<[^<>\r\n]+>>", line):
            malformed.append(line.strip()[:120])
    return {
        "row_fields": row_fields,
        "common_fields": common_fields,
        "conflicting_fields": sorted(set(row_fields) & set(common_fields)),
        "malformed_common_placeholders": sorted(set(x for x in malformed if x)),
    }
# =====================================================
# ファイル名生成
# =====================================================
def sanitize_filename_part(value):
    value = str(value).strip()
    value = re.sub(r'[\\/:*?"<>|]', "_", value)
    value = re.sub(r"\s+", " ", value)
    return value


def build_output_filename(template_path, row_dict, filename_keys=None, row_index=1,
                          add_serial_number=False, serial_digits=3, extension=None):
    template_name = os.path.basename(template_path)
    template_stem, template_ext = os.path.splitext(template_name)
    ext = extension or template_ext
    if not ext.startswith("."):
        ext = "." + ext
    parts = []
    if add_serial_number:
        parts.append(f"{row_index:0{serial_digits}d}")
    if filename_keys:
        for key in filename_keys:
            value = sanitize_filename_part(row_dict.get(key, ""))
            if value:
                parts.append(value)
    parts.append(sanitize_filename_part(template_stem))
    return "_".join(parts) + ext


def ensure_unique_path(path):
    if not os.path.exists(path):
        return path
    folder = os.path.dirname(path)
    name = os.path.basename(path)
    stem, ext = os.path.splitext(name)
    counter = 2
    while True:
        candidate = os.path.join(folder, f"{stem}_{counter}{ext}")
        if not os.path.exists(candidate):
            return candidate
        counter += 1


# =====================================================
# PDF変換・PDF分割
# =====================================================
def convert_docx_to_pdf_with_word(docx_path, pdf_path):
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"PDF変換元のWordファイルが見つかりません: {docx_path}")
    try:
        import pythoncom
        import win32com.client
    except Exception as e:
        raise RuntimeError("PDF出力には pywin32 が必要です。\n\npip install pywin32") from e

    abs_docx_path = os.path.abspath(docx_path)
    abs_pdf_path = os.path.abspath(pdf_path)
    word = None
    doc = None
    try:
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(abs_docx_path)
        doc.SaveAs2(abs_pdf_path, FileFormat=17)
    except Exception as e:
        raise RuntimeError(
            "Microsoft WordによるPDF変換に失敗しました。\n"
            "このPCにMicrosoft Wordがインストールされているか、対象ファイルをWordで開けるか確認してください。\n\n"
            f"詳細: {e}"
        ) from e
    finally:
        try:
            if doc is not None:
                doc.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def split_pdf_evenly_by_record_count(input_pdf_path, output_pdf_paths):
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(input_pdf_path)
    total_pages = len(reader.pages)
    record_count = len(output_pdf_paths)
    if record_count == 0:
        raise ValueError("分割対象のデータがありません。")
    if total_pages % record_count != 0:
        return False, total_pages, record_count, None
    pages_per_record = total_pages // record_count
    if pages_per_record <= 0:
        return False, total_pages, record_count, None

    for i, output_pdf_path in enumerate(output_pdf_paths):
        writer = PdfWriter()
        start_page = i * pages_per_record
        end_page = start_page + pages_per_record
        for page_index in range(start_page, end_page):
            writer.add_page(reader.pages[page_index])
        with open(output_pdf_path, "wb") as f:
            writer.write(f)
        writer.close()
    return True, total_pages, record_count, pages_per_record


# =====================================================
# Word/PDF生成
# =====================================================
def row_to_replacements(row):
    return {str(k): "" if v is None else str(v) for k, v in row.to_dict().items()}


def build_pdf_output_paths(template_path, df, filename_keys, output_dir_path, add_serial_number=False, serial_digits=3):
    paths = []
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        row_dict = row_to_replacements(row)
        pdf_filename = build_output_filename(
            template_path, row_dict, filename_keys, i,
            add_serial_number=add_serial_number, serial_digits=serial_digits, extension=".pdf"
        )
        paths.append(ensure_unique_path(os.path.join(output_dir_path, pdf_filename)))
    return paths


def generate_documents_from_template(template_path, df, filename_keys, output_dir_path,
                                     add_serial_number=False, serial_digits=3, common_replacements=None, progress_callback=None):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"テンプレートWordが見つかりません: {template_path}")
    if df is None or df.empty:
        raise ValueError("置換データがありません。")
    os.makedirs(output_dir_path, exist_ok=True)
    total = len(df)
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        row_dict = row_to_replacements(row)
        doc = Document(template_path)
        replace_placeholders(doc, row_dict, common_replacements)
        filename = build_output_filename(template_path, row_dict, filename_keys, i, add_serial_number, serial_digits, extension=".docx")
        output_path = ensure_unique_path(os.path.join(output_dir_path, filename))
        doc.save(output_path)
        if progress_callback:
            progress_callback(i, total)


def generate_pdf_documents_from_template(template_path, df, filename_keys, output_dir_path,
                                         add_serial_number=False, serial_digits=3, common_replacements=None, progress_callback=None):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"テンプレートWordが見つかりません: {template_path}")
    if df is None or df.empty:
        raise ValueError("置換データがありません。")
    os.makedirs(output_dir_path, exist_ok=True)
    total = len(df)
    with tempfile.TemporaryDirectory() as tmp_dir:
        for i, (_, row) in enumerate(df.iterrows(), start=1):
            row_dict = row_to_replacements(row)
            doc = Document(template_path)
            replace_placeholders(doc, row_dict, common_replacements)
            docx_filename = build_output_filename(template_path, row_dict, filename_keys, i, add_serial_number, serial_digits, extension=".docx")
            pdf_filename = os.path.splitext(docx_filename)[0] + ".pdf"
            temp_docx_path = os.path.join(tmp_dir, docx_filename)
            output_pdf_path = ensure_unique_path(os.path.join(output_dir_path, pdf_filename))
            doc.save(temp_docx_path)
            convert_docx_to_pdf_with_word(temp_docx_path, output_pdf_path)
            if progress_callback:
                progress_callback(i, total)


def iter_body_elements_without_sectpr(doc):
    body = doc._element.body
    for element in body:
        if element.tag.endswith('}sectPr'):
            continue
        yield element


def append_document_body(target_doc, source_doc):
    target_body = target_doc._element.body
    sect_pr = target_body.sectPr
    for element in iter_body_elements_without_sectpr(source_doc):
        new_element = copy.deepcopy(element)
        if sect_pr is not None:
            target_body.insert(target_body.index(sect_pr), new_element)
        else:
            target_body.append(new_element)


def generate_merged_document_from_template(template_path, df, output_docx_path,
                                          add_serial_number=False, serial_digits=3, common_replacements=None, progress_callback=None):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"テンプレートWordが見つかりません: {template_path}")
    if df is None or df.empty:
        raise ValueError("置換データがありません。")
    output_dir = os.path.dirname(output_docx_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
    total = len(df)
    merged_doc = None
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        row_dict = row_to_replacements(row)
        doc = Document(template_path)
        replace_placeholders(doc, row_dict, common_replacements)
        if merged_doc is None:
            merged_doc = doc
        else:
            merged_doc.add_page_break()
            append_document_body(merged_doc, doc)
        if progress_callback:
            progress_callback(i, total)
    merged_doc.save(output_docx_path)


def generate_merged_pdf_document_from_template(template_path, df, output_pdf_path,
                                             add_serial_number=False, serial_digits=3, common_replacements=None, progress_callback=None):
    if df is None or df.empty:
        raise ValueError("置換データがありません。")
    output_dir = os.path.dirname(output_pdf_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp_dir:
        temp_docx_path = os.path.join(tmp_dir, "merged_temp.docx")
        generate_merged_document_from_template(template_path, df, temp_docx_path,
                                             add_serial_number, serial_digits, common_replacements=common_replacements, progress_callback=progress_callback)
        convert_docx_to_pdf_with_word(temp_docx_path, output_pdf_path)


def generate_pdf_documents_fast_or_fallback(template_path, df, filename_keys, output_dir_path,
                                            add_serial_number=False, serial_digits=3,
                                            common_replacements=None, progress_callback=None, fallback_callback=None):
    if df is None or df.empty:
        raise ValueError("置換データがありません。")
    os.makedirs(output_dir_path, exist_ok=True)
    output_pdf_paths = build_pdf_output_paths(template_path, df, filename_keys, output_dir_path,
                                             add_serial_number, serial_digits)
    try:
        with tempfile.TemporaryDirectory() as tmp_dir:
            merged_docx_path = os.path.join(tmp_dir, "merged_temp.docx")
            merged_pdf_path = os.path.join(tmp_dir, "merged_temp.pdf")
            generate_merged_document_from_template(template_path, df, merged_docx_path,
                                                 add_serial_number, serial_digits, common_replacements=common_replacements, progress_callback=progress_callback)
            convert_docx_to_pdf_with_word(merged_docx_path, merged_pdf_path)
            ok, total_pages, record_count, pages_per_record = split_pdf_evenly_by_record_count(merged_pdf_path, output_pdf_paths)
            if ok:
                if progress_callback:
                    progress_callback(record_count, record_count)
                return {"mode": "fast", "total_pages": total_pages, "record_count": record_count, "pages_per_record": pages_per_record}
            if fallback_callback:
                fallback_callback(total_pages, record_count)
    except ImportError:
        if fallback_callback:
            fallback_callback(None, len(df))
    except Exception:
        if fallback_callback:
            fallback_callback(None, len(df))

    generate_pdf_documents_from_template(template_path, df, filename_keys, output_dir_path,
                                         add_serial_number, serial_digits, common_replacements=common_replacements, progress_callback=progress_callback)
    return {"mode": "fallback", "record_count": len(df)}


def generate_zip_from_template(template_path, df, filename_keys, output_zip_path,
                               add_serial_number=False, serial_digits=3, common_replacements=None, progress_callback=None):
    if df is None or df.empty:
        raise ValueError("置換データがありません。")
    output_dir = os.path.dirname(output_zip_path)
    os.makedirs(output_dir, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp_dir:
        generate_documents_from_template(template_path, df, filename_keys, tmp_dir,
                                         add_serial_number, serial_digits, common_replacements=common_replacements, progress_callback=progress_callback)
        with zipfile.ZipFile(output_zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
            for filename in os.listdir(tmp_dir):
                file_path = os.path.join(tmp_dir, filename)
                if os.path.isfile(file_path):
                    zipf.write(file_path, arcname=filename)


def generate_pdf_zip_from_template(template_path, df, filename_keys, output_zip_path,
                                   add_serial_number=False, serial_digits=3, common_replacements=None, progress_callback=None):
    if df is None or df.empty:
        raise ValueError("置換データがありません。")
    output_dir = os.path.dirname(output_zip_path)
    os.makedirs(output_dir, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp_dir:
        generate_pdf_documents_from_template(template_path, df, filename_keys, tmp_dir,
                                             add_serial_number, serial_digits, common_replacements=common_replacements, progress_callback=progress_callback)
        with zipfile.ZipFile(output_zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
            for filename in os.listdir(tmp_dir):
                file_path = os.path.join(tmp_dir, filename)
                if os.path.isfile(file_path):
                    zipf.write(file_path, arcname=filename)


def generate_pdf_zip_fast_or_fallback(template_path, df, filename_keys, output_zip_path,
                                      add_serial_number=False, serial_digits=3,
                                      common_replacements=None, progress_callback=None, fallback_callback=None):
    if df is None or df.empty:
        raise ValueError("置換データがありません。")
    output_dir = os.path.dirname(output_zip_path)
    os.makedirs(output_dir, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp_dir:
        generate_pdf_documents_fast_or_fallback(template_path, df, filename_keys, tmp_dir,
                                                add_serial_number, serial_digits, common_replacements=common_replacements,
                                                progress_callback=progress_callback, fallback_callback=fallback_callback)
        with zipfile.ZipFile(output_zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
            for filename in os.listdir(tmp_dir):
                file_path = os.path.join(tmp_dir, filename)
                if os.path.isfile(file_path):
                    zipf.write(file_path, arcname=filename)
